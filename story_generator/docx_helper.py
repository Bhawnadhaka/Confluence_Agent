# modules/story_generator/docx_helpers.py
from docx import Document

def _make_cell_bold(cell):
    """Make cell text bold"""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

def _merge_cells(table, start_row, start_col, end_row, end_col):
    """Merge table cells"""
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            if row == start_row and col == start_col:
                continue
            table.cell(row, col).merge(table.cell(start_row, start_col))
